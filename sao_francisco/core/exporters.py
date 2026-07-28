"""Deterministic TXT, DOCX, SRT and WebVTT exporters."""

from __future__ import annotations

import os
import re
import tempfile
from collections.abc import Callable
from pathlib import Path

from .models import Segment, Transcript


class ExportError(RuntimeError):
    """A transcript could not be exported."""


def format_timestamp(
    seconds: float,
    *,
    decimal_separator: str = ".",
    include_milliseconds: bool = True,
) -> str:
    """Format an unbounded media timestamp without wrapping after 24 hours."""

    if seconds < 0:
        raise ValueError("timestamp must not be negative")
    total_milliseconds = int(round(seconds * 1000))
    hours, remainder = divmod(total_milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    whole_seconds, milliseconds = divmod(remainder, 1_000)
    base = f"{hours:02d}:{minutes:02d}:{whole_seconds:02d}"
    if include_milliseconds:
        return f"{base}{decimal_separator}{milliseconds:03d}"
    return base


def _non_empty_segments(transcript: Transcript) -> tuple[Segment, ...]:
    return tuple(
        segment
        for segment in transcript.segments
        if segment.text.strip() and segment.end > segment.start
    )


def transcript_to_txt(
    transcript: Transcript,
    *,
    include_timestamps: bool = False,
    include_speakers: bool = True,
) -> str:
    if not include_timestamps:
        paragraphs = []
        for speaker, text in _natural_paragraphs(transcript):
            speaker_prefix = f"{speaker}: " if include_speakers and speaker else ""
            paragraphs.append(f"{speaker_prefix}{text}")
        return "\n\n".join(paragraphs).rstrip() + ("\n" if paragraphs else "")

    lines: list[str] = []
    for segment in transcript.segments:
        prefix: list[str] = []
        prefix.append(
            f"[{format_timestamp(segment.start, include_milliseconds=False)}]"
        )
        if include_speakers and segment.speaker:
            prefix.append(f"{segment.speaker}:")
        lines.append(" ".join((*prefix, segment.text)).strip())
    return "\n".join(lines).rstrip() + ("\n" if lines else "")


def _natural_paragraphs(transcript: Transcript) -> tuple[tuple[str | None, str], ...]:
    """Merge cue-sized segments into readable paragraphs without changing words."""

    soft_limit = 900
    hard_limit = 1_600
    source_groups: list[tuple[str | None, str]] = []
    current_speaker: str | None = None
    current_parts: list[str] = []
    previous_end: float | None = None

    def flush_source_group() -> None:
        nonlocal current_parts
        if current_parts:
            source_groups.append((current_speaker, " ".join(current_parts)))
        current_parts = []

    for segment in _non_empty_segments(transcript):
        text = segment.text.strip()
        speaker_changed = bool(current_parts) and segment.speaker != current_speaker
        long_pause = (
            bool(current_parts)
            and previous_end is not None
            and segment.start - previous_end >= 2.0
        )
        if speaker_changed or long_pause:
            flush_source_group()
        if not current_parts:
            current_speaker = segment.speaker
        current_parts.append(text)
        previous_end = segment.end
    flush_source_group()

    paragraphs: list[tuple[str | None, str]] = []
    for speaker, group_text in source_groups:
        sentences = re.split(r"(?<=[.!?…])\s+", group_text)
        packed: list[str] = []
        packed_length = 0
        for sentence in sentences:
            sentence = sentence.strip()
            if not sentence:
                continue
            next_length = packed_length + (1 if packed else 0) + len(sentence)
            if packed and packed_length >= soft_limit and next_length > hard_limit:
                paragraphs.append((speaker, " ".join(packed)))
                packed = []
                packed_length = 0
            packed.append(sentence)
            packed_length += (1 if packed_length else 0) + len(sentence)
        if packed:
            tail = " ".join(packed)
            if (
                len(tail) < 300
                and paragraphs
                and paragraphs[-1][0] == speaker
                and len(paragraphs[-1][1]) + 1 + len(tail) <= hard_limit + 300
            ):
                previous_speaker, previous_text = paragraphs[-1]
                paragraphs[-1] = (previous_speaker, f"{previous_text} {tail}")
            else:
                paragraphs.append((speaker, tail))
    return tuple(paragraphs)


def transcript_to_srt(transcript: Transcript) -> str:
    cues: list[str] = []
    for index, segment in enumerate(_non_empty_segments(transcript), start=1):
        timing = (
            f"{format_timestamp(segment.start, decimal_separator=',')} --> "
            f"{format_timestamp(segment.end, decimal_separator=',')}"
        )
        speaker = f"{segment.speaker}: " if segment.speaker else ""
        cues.append(f"{index}\n{timing}\n{speaker}{segment.text}")
    return "\n\n".join(cues).rstrip() + ("\n" if cues else "")


def transcript_to_vtt(transcript: Transcript) -> str:
    cues: list[str] = []
    for segment in _non_empty_segments(transcript):
        timing = (
            f"{format_timestamp(segment.start)} --> "
            f"{format_timestamp(segment.end)}"
        )
        text = (
            f"<v {segment.speaker}>{segment.text}</v>"
            if segment.speaker
            else segment.text
        )
        cues.append(f"{timing}\n{text}")
    payload = "\n\n".join(cues)
    return f"WEBVTT\n\n{payload}".rstrip() + "\n"


def _atomic_text(path: Path, content: str) -> Path:
    path = path.expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as stream:
            stream.write(content)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return path


def export_txt(
    transcript: Transcript,
    path: str | os.PathLike[str],
    *,
    include_timestamps: bool = False,
    include_speakers: bool = True,
) -> Path:
    return _atomic_text(
        Path(path),
        transcript_to_txt(
            transcript,
            include_timestamps=include_timestamps,
            include_speakers=include_speakers,
        ),
    )


def export_srt(transcript: Transcript, path: str | os.PathLike[str]) -> Path:
    return _atomic_text(Path(path), transcript_to_srt(transcript))


def export_vtt(transcript: Transcript, path: str | os.PathLike[str]) -> Path:
    return _atomic_text(Path(path), transcript_to_vtt(transcript))


def export_docx(
    transcript: Transcript,
    path: str | os.PathLike[str],
    *,
    include_timestamps: bool = False,
    include_speakers: bool = True,
    title: str | None = None,
) -> Path:
    """Write a simple, portable Word document using python-docx."""

    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as error:
        raise ExportError("DOCX export requires the 'python-docx' package") from error

    destination = Path(path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.", suffix=".docx", dir=destination.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document = Document()
        normal_style = document.styles["Normal"]
        normal_style.font.name = "Source Sans 3"
        normal_style.font.size = Pt(11)
        document.core_properties.language = transcript.language or ""
        document.core_properties.title = title or str(
            transcript.metadata.get("title", "")
        )
        if title:
            document.add_heading(title, level=0)

        if include_timestamps:
            for segment in transcript.segments:
                paragraph = document.add_paragraph()
                timestamp = format_timestamp(
                    segment.start, include_milliseconds=False
                )
                run = paragraph.add_run(f"[{timestamp}] ")
                run.bold = True
                if include_speakers and segment.speaker:
                    run = paragraph.add_run(f"{segment.speaker}: ")
                    run.bold = True
                paragraph.add_run(segment.text)
        else:
            for speaker, text in _natural_paragraphs(transcript):
                paragraph = document.add_paragraph()
                if include_speakers and speaker:
                    run = paragraph.add_run(f"{speaker}: ")
                    run.bold = True
                paragraph.add_run(text)
        document.save(str(temporary))
        with temporary.open("rb") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return destination


def export_improved_txt(text: str, path: str | os.PathLike[str]) -> Path:
    """Export already-paragraphed editorial text without timestamps."""

    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    return _atomic_text(Path(path), normalized + ("\n" if normalized else ""))


def export_improved_docx(
    text: str,
    path: str | os.PathLike[str],
    *,
    title: str | None = None,
    language: str | None = None,
) -> Path:
    """Export editorial paragraphs while preserving their existing order."""

    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as error:
        raise ExportError("DOCX export requires the 'python-docx' package") from error

    destination = Path(path).expanduser().resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{destination.stem}.", suffix=".docx", dir=destination.parent
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        document = Document()
        normal_style = document.styles["Normal"]
        normal_style.font.name = "Source Sans 3"
        normal_style.font.size = Pt(11)
        document.core_properties.language = language or ""
        document.core_properties.title = title or ""
        if title:
            document.add_heading(title, level=0)
        for value in re.split(r"\n{2,}", text.strip()):
            paragraph_text = value.strip()
            if paragraph_text:
                document.add_paragraph(paragraph_text)
        document.save(str(temporary))
        with temporary.open("rb") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return destination


_EXPORTERS: dict[str, Callable[..., Path]] = {
    "txt": export_txt,
    "srt": export_srt,
    "vtt": export_vtt,
    "docx": export_docx,
}


def export_transcript(
    transcript: Transcript,
    path: str | os.PathLike[str],
    *,
    format: str | None = None,
    **options: object,
) -> Path:
    """Export based on an explicit format or the destination suffix."""

    destination = Path(path)
    selected = (format or destination.suffix).lower().lstrip(".")
    try:
        exporter = _EXPORTERS[selected]
    except KeyError as error:
        raise ValueError(f"unsupported export format {selected!r}") from error
    return exporter(transcript, destination, **options)
